"""sample/給与計算.xlsx から計算結果を JSON fixture として抽出する.

サンプル人物データから計算結果のスナップショットを取り、Web 計算エンジンの
回帰テスト fixture として保存する。出力ファイルは gitignore 対象。
"""
import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3]
EXCEL = ROOT / "sample" / "給与計算.xlsx"
OUTPUT = Path(__file__).parent / "excel-snapshot.json"


def main() -> None:
    if not EXCEL.exists():
        raise SystemExit(f"Excel ファイルが見つかりません: {EXCEL}")

    wb = load_workbook(EXCEL, data_only=True)
    settings = wb["設定"]
    monthly = wb["月次計算"]

    birth_cell = settings["B3"].value
    birth = (
        birth_cell.strftime("%Y-%m-%d")
        if isinstance(birth_cell, date)
        else (birth_cell or "")
    )

    cases = []
    for row in monthly.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        year, month = int(row[0]), int(row[1])
        std = int(row[4])
        gross = int(row[5])
        cases.append(
            {
                "year": year,
                "month": month,
                "input": {
                    "stdRemuneration": std,
                    "grossSalary": gross,
                    "birthDate": birth,
                },
                # 列対応 (月次計算シート・告知書単位の合算丸めモデル):
                #   Q=協会けんぽ全額(17) R=協会社員(18) S=協会事業主(19)
                #   T=厚年全額(20) U=厚年社員(21) V=厚年事業主(22)
                #   W=拠出金全額(23) X=拠出金事業主(24)
                #   Y=社員天引(25) Z=事業主負担(26) AA=納付額(27) AB=差引支給(28)
                # values_only タプルは 0 始まりなので「列番号 − 1」で参照する。
                "expected": {
                    "isKaigoApplicable": bool(row[3]),
                    "kyokaiTotal": int(row[16]) if row[16] is not None else None,
                    "kyokaiEmployee": int(row[17]) if row[17] is not None else None,
                    "kyokaiEmployer": int(row[18]) if row[18] is not None else None,
                    "koseiTotal": int(row[19]) if row[19] is not None else None,
                    "koseiEmployee": int(row[20]) if row[20] is not None else None,
                    "koseiEmployer": int(row[21]) if row[21] is not None else None,
                    "kosodateTotal": int(row[22]) if row[22] is not None else None,
                    "kosodateEmployer": int(row[23]) if row[23] is not None else None,
                    "employeeDeductionTotal": int(row[24]) if row[24] is not None else None,
                    "employerBurdenTotal": int(row[25]) if row[25] is not None else None,
                    "payableTotal": int(row[26]) if row[26] is not None else None,
                    "netSalary": int(row[27]) if row[27] is not None else None,
                },
            }
        )

    OUTPUT.write_text(
        json.dumps(
            {
                "extractedAt": datetime.now().isoformat(),
                "sourceFile": str(EXCEL.relative_to(ROOT)),
                "birthDate": birth,
                "cases": cases,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    print(f"Wrote {OUTPUT} with {len(cases)} cases")


if __name__ == "__main__":
    main()
